import json, os, re, threading, time
from pathlib import Path
from flask import Flask, request, jsonify, send_from_directory
import ollama

ROOT = Path(__file__).resolve().parent
UPLOADS = ROOT / "uploads"
UPLOADS.mkdir(exist_ok=True)
MEMORY_FILE = ROOT / "memoria.json"
CATEGORIES_FILE = ROOT / "categorias_memoria.json"
MODEL = os.getenv("SALU_MODEL", "qwen3:1.7b")
MAX_FILE_MB = 20

app = Flask(__name__, static_folder="static")
app.config["MAX_CONTENT_LENGTH"] = MAX_FILE_MB * 1024 * 1024
lock = threading.Lock()
# Historial separado por ventana/sesión del navegador; no mezclar conversaciones.
histories = {}
documents = {}

def load_json(path, default):
    try:
        return json.loads(path.read_text(encoding="utf-8")) if path.exists() else default
    except (OSError, json.JSONDecodeError):
        return default

memory = load_json(MEMORY_FILE, {})
categories = load_json(CATEGORIES_FILE, {})
if isinstance(memory, list):
    memory = {"otros": memory}

def save_memory():
    MEMORY_FILE.write_text(json.dumps(memory, ensure_ascii=False, indent=2), encoding="utf-8")
    CATEGORIES_FILE.write_text(json.dumps(categories, ensure_ascii=False, indent=2), encoding="utf-8")

def extract_text(path):
    ext = path.suffix.lower()
    try:
        if ext == ".txt" or ext in {".py", ".js", ".html", ".css", ".json", ".md", ".csv", ".log"}:
            return path.read_text(encoding="utf-8", errors="replace")[:120000]
        if ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            return "\n\n".join(f"[Página {i+1}]\n{p.extract_text() or ''}" for i, p in enumerate(reader.pages))[:120000]
        if ext == ".docx":
            from docx import Document
            doc = Document(str(path))
            return "\n".join(p.text for p in doc.paragraphs)[:120000]
        if ext in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            chunks = []
            for ws in wb.worksheets:
                chunks.append(f"[Hoja: {ws.title}]")
                for row in ws.iter_rows(values_only=True):
                    vals = [str(v) for v in row if v is not None]
                    if vals: chunks.append(" | ".join(vals))
            return "\n".join(chunks)[:120000]
        if ext in {".png", ".jpg", ".jpeg", ".webp", ".bmp"}:
            try:
                import pytesseract
                from PIL import Image
                return "[Texto extraído de imagen mediante OCR]\n" + pytesseract.image_to_string(Image.open(path), lang="spa+eng")[:30000]
            except Exception:
                return "[Imagen recibida. Para analizar su contenido instala Tesseract OCR y pytesseract, o usa un modelo Ollama con visión.]"
    except Exception as e:
        return f"[No se pudo extraer texto: {e}]"
    return f"[Formato {ext} no compatible todavía]"

@app.get("/")
def index():
    return send_from_directory("static", "index.html")

@app.get("/api/status")
def status():
    try:
        tags = ollama.list()
        models = [m.get("name", m.get("model", "")) for m in tags.get("models", [])]
        return jsonify(ok=True, model=MODEL, ollama=True, models=models, memory_count=len(memory))
    except Exception as e:
        return jsonify(ok=False, model=MODEL, ollama=False, error=str(e), memory_count=len(memory))

@app.post("/api/upload")
def upload():
    if "file" not in request.files:
        return jsonify(error="No llegó ningún archivo."), 400
    f = request.files["file"]
    if not f.filename:
        return jsonify(error="El archivo no tiene nombre."), 400
    safe = re.sub(r"[^A-Za-z0-9._ -]", "_", Path(f.filename).name)[:120]
    path = UPLOADS / f"{int(time.time()*1000)}_{safe}"
    f.save(path)
    content = extract_text(path)
    doc_id = path.name
    documents[doc_id] = {"name": f.filename, "path": str(path), "text": content}
    return jsonify(id=doc_id, name=f.filename, chars=len(content), preview=content[:900], note=("No se extrajo texto legible; puede ser un PDF escaneado." if not content.strip() else ""))


def remember_from_message(prompt):
    """Guarda de forma explícita datos personales sencillos que el usuario pide recordar."""
    text = prompt.strip()
    patterns = [
        (r"(?i)\bmi\s+anime\s+favorito\s+es\s+(.+?)[.!]?$", "anime favorito", "preferencias"),
        (r"(?i)\bmi\s+comida\s+favorita\s+es\s+(.+?)[.!]?$", "comida favorita", "preferencias"),
        (r"(?i)\bmi\s+color\s+favorito\s+es\s+(.+?)[.!]?$", "color favorito", "preferencias"),
        (r"(?i)\bmi\s+nombre\s+es\s+(.+?)[.!]?$", "nombre", "perfil"),
        (r"(?i)\brecuerda\s+que\s+(.+?)[.!]?$", None, "otros"),
        (r"(?i)\brecuerda\s+(.+?)[.!]?$", None, "otros"),
        (r"(?i)\bguarda\s+que\s+(.+?)[.!]?$", None, "otros"),
    ]
    for pattern, fixed_key, category in patterns:
        match = re.search(pattern, text)
        if not match:
            continue
        value = match.group(1).strip().strip('"“”')
        if not value:
            continue
        if fixed_key:
            key = fixed_key
        else:
            # Para frases libres, guardar la afirmación completa como recuerdo.
            key = value[:100]
        memory[key] = value if fixed_key else "El usuario pidió recordar: " + value
        categories[key] = category
        save_memory()
        return key
    return None

def memory_question(prompt):
    """Responde consultas directas usando la memoria local antes de llamar al modelo."""
    p = prompt.strip().lower()
    aliases = {
        "anime favorito": ["¿Cuál es mi anime favorito?", "cual es mi anime favorito", "qué anime me gusta más", "que anime me gusta mas"],
        "comida favorita": ["¿Cuál es mi comida favorita?", "cual es mi comida favorita"],
        "color favorito": ["¿Cuál es mi color favorito?", "cual es mi color favorito"],
        "nombre": ["¿Cómo me llamo?", "como me llamo", "cuál es mi nombre", "cual es mi nombre"],
    }
    for key, questions in aliases.items():
        if key in memory and any(q.lower().strip("¿?") in p.strip("¿?") or p.strip("¿?") in q.lower().strip("¿?") for q in questions):
            return f"Según mi memoria, tu {key} es {memory[key]}."
    if any(phrase in p for phrase in ("ver memoria", "mostrar memoria", "qué recuerdas de mí", "que recuerdas de mi", "qué sabes de mí", "que sabes de mi")):
        if not memory:
            return "Todavía no hay recuerdos guardados."
        return "## Memoria de SaLu\n" + "\n".join(f"- **{k}**: {v}" for k, v in memory.items())
    return None

@app.post("/api/chat")
def chat():
    data = request.get_json(force=True)
    prompt = (data.get("message") or "").strip()
    selected = data.get("documents") or []
    session_id = re.sub(r"[^A-Za-z0-9_-]", "", str(data.get("session_id", "default")))[:100] or "default"
    if not prompt:
        return jsonify(error="Escribe un mensaje."), 400
    direct_answer = memory_question(prompt)
    if direct_answer:
        return jsonify(answer=direct_answer, model=MODEL, memory_count=len(memory))
    remembered_key = remember_from_message(prompt)
    if remembered_key:
        return jsonify(answer=f"¡Listo! Guardé en mi memoria: **{remembered_key}**. Puedes pedirme que lo recuerde más adelante.", model=MODEL, memory_count=len(memory))
    attached = []
    for doc_id in selected:
        doc = documents.get(doc_id)
        if doc:
            attached.append(f"DOCUMENTO: {doc['name']}\n{doc['text'][:24000]}")
    system = (
        "Eres SaLu, un asistente local cuyo idioma predeterminado es el español. "
        "RESPONDE SIEMPRE EN ESPAÑOL, incluso si el historial, una cita o un documento está en inglés, "
        "salvo que el usuario pida explícitamente otro idioma. "
        "Contesta la pregunta actual de forma directa; no respondas una pregunta anterior ni cambies de tema. "
        "Usa el historial solo si ayuda a entender referencias como 'ese ejercicio' o 'lo anterior'. "
        "Si no tienes suficiente contexto para identificar a qué se refiere, pregunta cuál ejercicio o tema indica. "
        "No inventes datos, nombres, citas ni hechos. Si no estás seguro, dilo y separa hechos de suposiciones. "
        "En temas especializados, evita afirmar detalles dudosos como si fueran ciertos. "
        "Usa Markdown y escribe fórmulas matemáticas en LaTeX delimitado por $...$ o $$...$$. "
        "Cuando el usuario adjunte documentos, basa el análisis en su contenido y señala si algo no aparece. "
        "No afirmes haber visto partes que no se hayan extraído."
    )
    mem_context = "\n".join(f"- [{categories.get(k,'perfil')}] {k}: {v}" for k,v in memory.items())
    if mem_context:
        system += "\nMemoria conocida del usuario:\n" + mem_context[:5000]
    msgs = [{"role":"system","content":system}]
    if attached:
        msgs.append({"role":"system","content":"CONTENIDO DE ARCHIVOS ADJUNTOS:\n\n" + "\n\n---\n\n".join(attached)})
    history = histories.setdefault(session_id, [])
    msgs.extend(history[-8:])
    msgs.append({"role":"user","content":prompt})
    try:
        with lock:
            result = ollama.chat(model=MODEL, messages=msgs)
        answer = result["message"]["content"]
        history.extend([{"role":"user","content":prompt},{"role":"assistant","content":answer}])
        del history[:-16]
        # Simple explicit memory commands; user can ask to see current memory.
        if prompt.lower() in {"ver memoria", "mostrar memoria"}:
            answer = "## Memoria de SaLu\n" + ("\n".join(f"- **{k}**: {v}" for k,v in memory.items()) if memory else "Todavía no hay recuerdos guardados.")
        return jsonify(answer=answer, model=MODEL)
    except Exception as e:
        return jsonify(error=f"No pude conectar con Ollama ({MODEL}). Comprueba que Ollama esté abierto y que el modelo esté instalado. Detalle: {e}"), 503

@app.post("/api/history/clear")
def clear_history():
    data = request.get_json(silent=True) or {}
    session_id = re.sub(r"[^A-Za-z0-9_-]", "", str(data.get("session_id", "default")))[:100] or "default"
    histories.pop(session_id, None)
    return jsonify(ok=True, message="Historial de conversación reiniciado.")


@app.get("/api/memory")
def get_memory():
    return jsonify(memory=memory, categories=categories)

@app.post("/api/memory")
def put_memory():
    data = request.get_json(force=True)
    key, value = str(data.get("key","")).strip(), str(data.get("value","")).strip()
    category = str(data.get("category","perfil")).strip() or "perfil"
    if not key or not value: return jsonify(error="Completa clave y valor."), 400
    memory[key] = value; categories[key] = category; save_memory()
    return jsonify(ok=True, memory=memory, categories=categories)

@app.delete("/api/memory/<path:key>")
def delete_memory(key):
    memory.pop(key, None); categories.pop(key, None); save_memory()
    return jsonify(ok=True)

@app.get("/api/documents")
def list_documents():
    return jsonify(documents=[{"id":k,"name":v["name"],"chars":len(v["text"])} for k,v in documents.items()])

if __name__ == "__main__":
    print(f"SaLu Web disponible en http://127.0.0.1:5000 | modelo: {MODEL}")
    app.run(host="127.0.0.1", port=5000, debug=False)
